"""Importa o "CRM Revendas & Eventos" (planilha Google do Apps Script) para uma empresa do SAX.

Lê o .xlsx baixado do Google Planilhas (ou o texto exportado com "## Sheet name:") e traz:
- LOJAS -> Unidades ("Curitiba - Batel"), USUARIOS -> Equipe com acesso por e-mail;
- FUNIS -> funis com as mesmas etapas; LEADS REVENDA/EVENTOS -> Leads e clientes;
- NEGOCIOS -> Negócios na etapa, com responsável, loja, supervisão e temperatura como etiquetas;
- ATIVIDADES -> Agenda; HISTORICO -> histórico do negócio.

Senhas NÃO são importadas (nem o hash da planilha): cada pessoa recebe um link de ativação e cria a sua.
Cada registro guarda o ID da planilha (`planilha_id`); importar de novo só traz o que ainda não veio.
"""

from __future__ import annotations

import csv
import io
import json
import re
import secrets
import unicodedata
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

from lib.db import client, controle
from models.common import new_id, now_utc

FUSO = ZoneInfo("America/Sao_Paulo")
PRACAS = {"CURITIBA": "Curitiba", "SAO PAULO": "São Paulo", "BRASILIA": "Brasília"}
PERFIS = {  # perfil da planilha -> (papel no SAX, cargo)
    "diretoria": ("admin", "Diretoria"),
    "gerente_op": ("admin", "Gerente de Operações"),
    "gerente_exp": ("admin", "Gerente de Expansão"),
    "supervisor": ("corretor", "Supervisor de Lojas"),
    "lider": ("corretor", "Líder de Loja"),
}
TIPOS_ATIVIDADE = {"contato": "ligacao", "ligacao": "ligacao", "whatsapp": "whatsapp", "e-mail": "email", "email": "email",
                   "reuniao com proprietario": "reuniao", "reuniao": "reuniao", "degustacao": "reuniao", "visita": "visita"}
REQUISITOS = {"contato": "contato", "reuniao com proprietario": "reuniao", "reuniao": "reuniao", "degustacao": "reuniao"}
FEITAS = {"sucesso", "concluida", "concluído", "concluido", "feito", "feita", "realizada", "ok"}


class ErroPlanilha(Exception):
    pass


def _sem_acento(t: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", t or "") if unicodedata.category(c) != "Mn").upper().strip()


def _praca(p: str | None) -> str | None:
    return PRACAS.get(_sem_acento(p or ""))


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M:%S")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def ler(conteudo: bytes) -> dict[str, list[list[str]]]:
    """Abas -> linhas (texto). Aceita .xlsx e o texto exportado com "## Sheet name:"."""
    if conteudo[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception:
            raise ErroPlanilha("Não consegui abrir o arquivo. Baixe a planilha em Arquivo > Fazer download > Microsoft Excel (.xlsx).")
        abas = {ws.title.strip().upper(): [[_txt(c) for c in linha] for linha in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
        wb.close()
        return abas
    try:
        texto = conteudo.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise ErroPlanilha("Formato não reconhecido. Envie o .xlsx baixado do Google Planilhas.")
    if "## Sheet name:" not in texto:
        raise ErroPlanilha("Formato não reconhecido. Envie o .xlsx baixado do Google Planilhas.")
    abas: dict[str, list[list[str]]] = {}
    partes = re.split(r"^## Sheet name: *(.+)$", texto, flags=re.M)
    for i in range(1, len(partes), 2):
        abas[partes[i].strip().upper()] = [[c.strip() for c in linha] for linha in csv.reader(io.StringIO(partes[i + 1].strip("\n")))]
    return abas


def _tabela(abas: dict, nome: str) -> list[dict]:
    linhas = [l for l in abas.get(nome, []) if any(c for c in l)]
    if not linhas:
        return []
    cab = [c.strip() for c in linhas[0]]
    saida = []
    for l in linhas[1:]:
        d = {cab[i]: (l[i] if i < len(l) else "") for i in range(len(cab)) if cab[i]}
        d["_extra"] = [c for i, c in enumerate(l) if c and (i >= len(cab) or not cab[i])]
        saida.append(d)
    return saida


def _data(v: str | None) -> datetime | None:
    v = (v or "").strip()
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(v, fmt).replace(tzinfo=FUSO).astimezone(timezone.utc)
        except ValueError:
            continue
    return None


def _dia(v: str | None) -> str | None:
    d = _data(v)
    return d.astimezone(FUSO).date().isoformat() if d else None


def _digitos(v: str | None) -> str:
    return "".join(c for c in (v or "") if c.isdigit())


def _valor(v: str | None) -> float | None:
    t = (v or "").replace("R$", "").strip()
    if not t:
        return None
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        n = float(t)
    except ValueError:
        return None
    return n if n > 0 else None


def _loja_curta(nome: str) -> str:
    return re.sub(r"^O Pudim Perfeito\s*-\s*", "", nome or "", flags=re.I).strip() or nome


async def _em_paralelo(tarefas: list, limite: int = 12) -> None:
    import asyncio

    sem = asyncio.Semaphore(limite)

    async def um(t):
        async with sem:
            await t

    await asyncio.gather(*(um(t) for t in tarefas))


def _rotear_local(pixel: str | None, supervisoes: list[dict]) -> dict | None:
    p = _sem_acento(pixel or "")
    if not p:
        return None
    for s in supervisoes:
        if s.get("ativa", True) and (p == _sem_acento(s.get("nome") or "") or p in {_sem_acento(x) for x in s.get("pixels") or []}):
            return s
    return None


async def importar(empresa: dict, conteudo: bytes, *, simular: bool, por: str, novos_links: bool = False) -> dict:
    """Rápido o bastante para o limite de 60 s do Vercel: lê o que já existe de uma vez e grava em lote.
    O que já veio antes (mesmo `planilha_id`) não duplica; ganha praça, supervisão e loja se faltava."""
    abas = ler(conteudo)
    if not any(k in abas for k in ("NEGOCIOS", "LEADS REVENDA", "USUARIOS")):
        raise ErroPlanilha("Esta planilha não tem as abas do CRM Revendas (USUARIOS, LEADS REVENDA, NEGOCIOS).")
    from lib.auth import hash_senha
    from lib.db import definir_empresa
    from lib.invites import invite
    from lib.tenant import indexar_usuario
    from models.pessoas import Pessoa, formatar_documento
    from models.usuarios import Usuario

    banco = client[empresa["db_name"]]
    definir_empresa(empresa["db_name"])
    agora = now_utc()
    cont = {k: 0 for k in ("lojas", "supervisoes", "usuarios", "clientes", "leads", "negocios", "atividades", "historico", "pedidos",
                           "ja_importados", "atualizados", "links_novos", "prints_pendentes")}
    avisos: list[str] = []
    acessos: list[dict] = []
    atualizar: list = []  # update_one pendentes (rodam em paralelo no fim)

    def mudar(colecao: str, doc: dict, campos: dict) -> None:
        novos = {k: v for k, v in campos.items() if v and doc.get(k) != v and not doc.get(k)}
        if novos:
            cont["atualizados"] += 1
            if not simular:
                atualizar.append(banco[colecao].update_one({"id": doc["id"]}, {"$set": novos}))

    # ------------------------------------------------------------ o que já existe (uma leitura por coleção)
    unidades = await banco.unidades.find({}).to_list(1000)
    sups_db = await banco.supervisoes.find({}).to_list(500)
    usuarios_db = {u["email"]: u for u in await banco.usuarios.find({}).to_list(2000) if u.get("email")}
    pessoas_db = await banco.pessoas.find({}, {"id": 1, "nome": 1, "telefone_digitos": 1, "planilha_id": 1, "papeis": 1}).to_list(50000)
    entradas_db = {d["planilha_id"]: d for d in await banco.entradas.find({"planilha_id": {"$ne": None}}).to_list(50000) if d.get("planilha_id")}
    leads_db = {d["planilha_id"]: d for d in await banco.leads.find({"planilha_id": {"$ne": None}}).to_list(50000) if d.get("planilha_id")}
    ativ_db = {d["planilha_id"]: d for d in await banco.atividades.find({"planilha_id": {"$ne": None}}, {"id": 1, "planilha_id": 1, "assunto": 1, "negocio_id": 1, "corretor_id": 1}).to_list(50000) if d.get("planilha_id")}
    # prints já trazidos ou à espera do arquivo (chave: ID do arquivo no Google Drive)
    drive_ja = {d["drive_id"] for d in await banco.anexos.find({"drive_id": {"$ne": None}}, {"drive_id": 1}).to_list(50000) if d.get("drive_id")}
    drive_ja |= {d["drive_id"] for d in await banco.anexos_pendentes.find({}, {"drive_id": 1}).to_list(50000) if d.get("drive_id")}
    hist_db = {d["planilha_id"] for d in await banco.historico.find({"planilha_id": {"$ne": None}}, {"planilha_id": 1}).to_list(100000) if d.get("planilha_id")}
    cfg = await banco.configuracoes.find_one({"id": "singleton"}) or {}

    # ------------------------------------------------------------ supervisões (SUP-... -> id)
    sup_map: dict[str, dict] = {}
    novas_sups = []
    for l in _tabela(abas, "SUPERVISOES"):
        sid, nome = l.get("ID") or "", (l.get("Supervisão") or "").strip()
        if not sid or not nome:
            continue
        praca = _praca(l.get("Praça")) or (l.get("Praça") or "").strip().title() or None
        pixels = sorted({p.strip() for p in (l.get("Pixels (separe por vírgula)") or "").split(",") if p.strip()})
        atual = next((x for x in sups_db if x.get("planilha_id") == sid or _sem_acento(x.get("nome") or "") == _sem_acento(nome)), None)
        if atual:
            sup_map[sid] = atual
            mudar("supervisoes", atual, {"praca": praca, "planilha_id": sid})
            continue
        doc = {"id": new_id(), "nome": nome, "praca": praca, "pixels": pixels, "ativa": (l.get("Ativa") or "SIM").upper() != "NÃO",
               "planilha_id": sid, "created_at": agora}
        sup_map[sid] = doc
        sups_db.append(doc)
        novas_sups.append(doc)
        cont["supervisoes"] += 1
    if novas_sups and not simular:
        await banco.supervisoes.insert_many(novas_sups)

    # ------------------------------------------------------------ lojas ("Curitiba - Batel")
    lojas: dict[str, dict] = {}
    linhas_lojas = _tabela(abas, "LOJAS")
    por_nome = {u.get("nome"): u for u in unidades}
    so_a_padrao = len(unidades) == 1 and unidades[0].get("nome") == cfg.get("nome_software")
    # supervisão de cada loja: a dos líderes daquela loja
    sup_da_loja: dict[str, str] = {}
    for l in _tabela(abas, "USUARIOS"):
        if l.get("Loja"):
            for s in (l.get("Supervisões") or "").split(","):
                if s.strip() in sup_map:
                    sup_da_loja.setdefault(l["Loja"], sup_map[s.strip()]["id"])
    novas_lojas = []
    for l in linhas_lojas:
        original = l.get("Loja") or ""
        if not original:
            continue
        cidade = _praca(l.get("Praça"))
        nome = f"{cidade} - {_loja_curta(original)}" if cidade else _loja_curta(original)
        u = por_nome.get(nome) or next((x for x in unidades if x.get("planilha_id") == original), None)
        if u:
            lojas[original] = u
            mudar("unidades", u, {"cidade": cidade, "supervisao_id": sup_da_loja.get(original), "planilha_id": original})
            continue
        doc = {"id": new_id(), "nome": nome, "principal": False, "ativa": True, "endereco": None, "cidade": cidade, "telefone": None,
               "whatsapp": None, "supervisao_id": sup_da_loja.get(original), "planilha_id": original, "created_at": agora}
        lojas[original] = doc
        novas_lojas.append(doc)
        cont["lojas"] += 1
    if novas_lojas and not simular:
        await banco.unidades.insert_many(novas_lojas)
        if so_a_padrao:  # a unidade automática da empresa sai; a primeira loja vira a principal
            await banco.unidades.delete_one({"id": unidades[0]["id"]})
            await banco.unidades.update_one({"id": novas_lojas[0]["id"]}, {"$set": {"principal": True}})

    # ------------------------------------------------------------ equipe e acessos (sem senha: link de ativação)
    pessoa_do_usuario: dict[str, str] = {}
    pessoa_por_nome = {_sem_acento(u.get("nome") or ""): u.get("pessoa_id") for u in usuarios_db.values() if u.get("pessoa_id")}
    novos_usuarios: list[tuple[dict, dict, dict]] = []
    convites: list[tuple[dict, dict]] = []
    for l in _tabela(abas, "USUARIOS"):
        uid_planilha, nome, email = l.get("ID") or "", (l.get("Nome") or "").strip(), (l.get("E-mail") or "").strip().lower()
        if not nome or "@" not in email:
            continue
        perfil = (l.get("Perfil") or "").strip().lower()
        papel, cargo = PERFIS.get(perfil, ("corretor", l.get("Perfil") or "Equipe"))
        praca_txt = (l.get("Praça") or "").strip()
        pracas = ["TODAS"] if _sem_acento(praca_txt) == "TODAS" else ([_praca(praca_txt)] if _praca(praca_txt) else [])
        sups = [sup_map[s.strip()]["id"] for s in (l.get("Supervisões") or "").split(",") if s.strip() in sup_map]
        loja = lojas.get(l.get("Loja") or "")
        nomes_sup = [sup_map[s.strip()]["nome"] for s in (l.get("Supervisões") or "").split(",") if s.strip() in sup_map]
        detalhe = " · ".join(x for x in [_loja_curta(l.get("Loja") or "") if loja else "", *nomes_sup] if x) or \
            ("Todas as praças" if pracas == ["TODAS"] else (pracas[0] if pracas else ""))
        cargo_final = f"{cargo} - {detalhe}" if detalhe and papel == "corretor" else cargo
        atuacao = {"perfil": perfil if perfil in PERFIS else None, "pracas": pracas, "supervisoes": sups, "unidade_id": (loja or {}).get("id"),
                   "planilha_id": uid_planilha}
        ativo = (l.get("Ativo") or "SIM").strip().upper() != "NÃO"
        item = {"nome": nome, "email": email, "perfil": cargo_final, "papel": "Gestor" if papel == "admin" else "Equipe", "link": None,
                "telefone": l.get("Telefone") or None}
        existente = usuarios_db.get(email)
        if existente:
            pessoa_do_usuario[uid_planilha] = existente.get("pessoa_id")
            mudar("usuarios", existente, {**atuacao, "cargo": cargo_final})
            if existente.get("activation_required") and existente.get("ativo", True):
                if novos_links:
                    item["situacao"] = "Ainda não ativou: link novo gerado (o anterior deixa de valer)"
                    convites.append((existente, item))
                else:
                    item["situacao"] = "Ainda não ativou"
            else:
                item["situacao"] = "Já tem acesso"
            acessos.append(item)
            continue
        if await controle.usuarios.find_one({"email": email}) or \
                any(i.get("empresa_id") != empresa["id"] for i in await controle.usuarios_index.find({"email": email}).to_list(5)):
            item["situacao"] = "E-mail já usado em outra empresa do SAX: não criado"
            acessos.append(item)
            avisos.append(f"{nome}: o e-mail {email} já é usado em outra empresa e não foi criado.")
            continue
        cont["usuarios"] += 1
        pessoa = Pessoa(nome=nome, papeis=["corretor"], email=email, telefone=l.get("Telefone") or None, cidade=pracas[0] if pracas and pracas[0] != "TODAS" else None).model_dump()
        pessoa["planilha_id"] = uid_planilha
        pessoa_do_usuario[uid_planilha] = pessoa["id"]
        pessoa_por_nome[_sem_acento(nome)] = pessoa["id"]
        usuario = Usuario(nome=nome, email=email, senha_hash=hash_senha(secrets.token_urlsafe(24)), papel=papel, pessoa_id=pessoa["id"],
                          telefone=l.get("Telefone") or None, cargo=cargo_final, ativo=ativo).model_dump()
        usuario.update(atuacao, activation_required=True)
        item["situacao"] = "Vai receber link para criar a senha" if ativo else "Inativo na planilha: criado sem acesso"
        novos_usuarios.append((pessoa, usuario, item))
        if ativo:
            convites.append((usuario, item))
        acessos.append(item)
    if novos_usuarios and not simular:
        await banco.pessoas.insert_many([p for p, _, _ in novos_usuarios])
        await banco.usuarios.insert_many([u for _, u, _ in novos_usuarios])
        for _, u, _ in novos_usuarios:
            await indexar_usuario(u["email"], empresa["id"], empresa["db_name"])
    if not simular:
        for u, item in convites:
            try:
                item["link"] = await invite(banco, u, empresa.get("nome") or "SAX")
                cont["links_novos"] += 1
            except Exception:
                item["situacao"] = "Criado; gere o link em Equipe > Link de acesso"

    # ------------------------------------------------------------ funis (etapas e requisitos da planilha)
    funis_planilha: dict[str, dict] = {}
    funis_db = await banco.funis.find({}).to_list(50)
    for l in _tabela(abas, "FUNIS"):
        fid, nome = l.get("ID") or "", (l.get("Nome") or "").strip()
        try:
            etapas = json.loads(l.get("Etapas (JSON)") or "[]")
        except ValueError:
            etapas = []
        if not fid or not nome or not etapas:
            continue
        atual = next((f for f in funis_db if f.get("planilha_id") == fid), None)
        if not atual:
            mesmo_nome = next((f for f in funis_db if (f.get("nome") or "").lower() == nome.lower()), None) or \
                next((f for f in funis_db if not f.get("planilha_id") and (f.get("nome") or "").lower().startswith(nome.lower())), None)
            vazio = mesmo_nome and not await banco.leads.find_one({"funil_id": mesmo_nome["id"]}, {"id": 1})
            n = len(etapas)
            novas = [{"id": new_id(), "nome": e.get("nome") or f"Etapa {i + 1}", "probabilidade": min(90, 10 + round(80 * i / max(n - 1, 1))),
                      "dias_parado": int(e.get("dias") or 0) or None, "cor": None, "planilha_id": e.get("id"),
                      "requisito": REQUISITOS.get(_sem_acento(e.get("requisito") or "").lower())} for i, e in enumerate(etapas)]
            if vazio:
                atual = {**mesmo_nome, "nome": nome, "etapas": novas, "planilha_id": fid}
                if not simular:
                    await banco.funis.update_one({"id": mesmo_nome["id"]}, {"$set": {"nome": nome, "etapas": novas, "planilha_id": fid}})
            else:
                atual = {"id": new_id(), "nome": nome if not mesmo_nome else f"{nome} (planilha)", "ordem": len(funis_db), "padrao": False,
                         "etapas": novas, "planilha_id": fid, "created_at": agora}
                if not simular:
                    await banco.funis.insert_one(atual)
                funis_db.append(atual)
        elif not simular:
            # funil já importado: garante os requisitos das etapas (versões antigas do importador não traziam)
            req = {e.get("id"): REQUISITOS.get(_sem_acento(e.get("requisito") or "").lower()) for e in etapas}
            etapas_db = [{**e, "requisito": e.get("requisito") or req.get(e.get("planilha_id"))} for e in atual["etapas"]]
            if etapas_db != atual["etapas"]:
                atual["etapas"] = etapas_db
                await banco.funis.update_one({"id": atual["id"]}, {"$set": {"etapas": etapas_db}})
        funis_planilha[fid] = {"funil": atual, "etapas": {e.get("planilha_id"): e for e in atual["etapas"]}}

    # ------------------------------------------------------------ leads e clientes
    negocios_linhas = _tabela(abas, "NEGOCIOS")
    negocio_do_lead = {n.get("ID Lead"): n for n in negocios_linhas if n.get("ID Lead")}
    tel_para_pessoa = {p["telefone_digitos"]: p for p in pessoas_db if p.get("telefone_digitos")}
    pessoa_por_planilha = {p["planilha_id"]: p for p in pessoas_db if p.get("planilha_id")}
    origens: set[str] = set()
    etiquetas: set[str] = set()
    cliente_do_lead: dict[str, str] = {}
    lead_info: dict[str, dict] = {}
    novos_clientes, novas_entradas = [], []
    negocio_novo_id: dict[str, str] = {n["ID"]: (leads_db.get(n["ID"]) or {}).get("id") or new_id() for n in negocios_linhas if n.get("ID")}

    def onde(neg: dict | None, pixel: str | None, localidade: str | None) -> dict:
        """praça, supervisão e loja: o negócio da planilha manda; senão o pixel; senão a localidade."""
        neg = neg or {}
        sup = sup_map.get(neg.get("Supervisão") or "") or _rotear_local(pixel, sups_db)
        loja = lojas.get(neg.get("Loja") or "")
        praca = _praca(neg.get("Praça")) or (sup or {}).get("praca") or (loja or {}).get("cidade") or _praca(localidade)
        return {"praca": praca, "supervisao_id": (sup or {}).get("id") or (loja or {}).get("supervisao_id"), "unidade_id": (loja or {}).get("id")}

    for aba, tipo in (("LEADS REVENDA", "Revenda"), ("LEADS EVENTOS", "Evento")):
        for l in _tabela(abas, aba):
            lid = l.get("ID Lead") or ""
            if not lid:
                continue
            restaurante = (l.get("Nome do Restaurante") or "").strip()
            responsavel = (l.get("Nome do Responsável") or "").strip()
            responsavel = "" if responsavel.lower() in ("sem nome", "-") else responsavel
            nome = restaurante or responsavel or "Contato da planilha"
            neg = negocio_do_lead.get(lid)
            local = onde(neg, l.get("Pixel"), l.get("Localidade do pedido"))
            lead_info[lid] = {"nome": nome, "tipo": tipo, "origem": (l.get("Origem") or "Planilha").strip()[:80], "pixel": (l.get("Pixel") or "").strip() or None, **local}
            origens.add(lead_info[lid]["origem"])
            existente = entradas_db.get(lid)
            if existente:
                cont["ja_importados"] += 1
                cliente_do_lead[lid] = existente.get("cliente_id")
                mudar("entradas", existente, {**local, "pixel": lead_info[lid]["pixel"]})
                continue
            tel = l.get("Telefone") or ""
            digitos = _digitos(tel)
            cliente = pessoa_por_planilha.get(lid) or (tel_para_pessoa.get(digitos) if len(digitos) >= 10 else None)
            if not cliente:
                cnpj = _digitos(l.get("CNPJ"))
                cnpj = "" if not cnpj.strip("0") else cnpj
                cliente = Pessoa(nome=nome, papeis=["cliente"], tipo_pessoa="pj" if (cnpj or (restaurante and tipo == "Revenda")) else "pf",
                                 nome_fantasia=restaurante or None, responsavel_nome=responsavel or None, telefone=tel or None,
                                 email=(l.get("E-mail") or "").strip() or None, cpf_cnpj=formatar_documento(cnpj) if cnpj else None,
                                 logradouro=(l.get("Endereço / bairro") or "")[:300] or None, cidade=l.get("Localidade do pedido") or None,
                                 observacoes=(l.get("Observações") or "")[:4000] or None).model_dump()
                cliente["telefone_digitos"] = digitos or None
                cliente["documento_digitos"] = cnpj or None
                cliente["planilha_id"] = lid
                cliente["created_at"] = _data(l.get("Data de entrada")) or agora
                novos_clientes.append(cliente)
                if digitos:
                    tel_para_pessoa[digitos] = cliente
                cont["clientes"] += 1
            cliente_do_lead[lid] = cliente["id"]
            dono = pessoa_do_usuario.get((neg or {}).get("Responsável ID") or "")
            criado = _data(l.get("Data de entrada")) or agora
            partes = [f"{tipo}: {restaurante}" if restaurante else tipo, f"Responsável: {responsavel}" if responsavel else "",
                      f"CNPJ: {l.get('CNPJ')}" if l.get("CNPJ") else "", f"Endereço: {l.get('Endereço / bairro')}" if l.get("Endereço / bairro") else "",
                      l.get("Observações") or "", f"Pixel: {l.get('Pixel')}" if l.get("Pixel") else "", *l.get("_extra", [])]
            novas_entradas.append({
                "id": new_id(), "nome": nome[:200], "telefone": tel or None, "email": (l.get("E-mail") or "").strip() or None,
                "origem": lead_info[lid]["origem"], "interesse": "outro", "mensagem": "\n".join(p for p in partes if p)[:2000] or None,
                "imovel_id": None, "veiculo_id": None, "cliente_id": cliente["id"], "corretor_id": dono, "valor_estimado": None,
                "etiquetas": [tipo], "status": "convertido" if neg else "novo", "negocio_id": negocio_novo_id.get((neg or {}).get("ID")) if neg else None,
                "motivo_descarte": None, "primeiro_contato_em": criado if neg else None, "atribuido_em": criado if dono else None,
                "na_fila": False, "fila_desde": None, "devolvido_de": None, "portal_lead_id": None, "pixel": lead_info[lid]["pixel"], **local,
                "planilha_id": lid, "created_at": criado, "updated_at": criado,
            })
            cont["leads"] += 1
    if not simular:
        if novos_clientes:
            await banco.pessoas.insert_many(novos_clientes)
        if novas_entradas:
            await banco.entradas.insert_many(novas_entradas)

    # ------------------------------------------------------------ negócios
    novos_negocios, novos_hist = [], []
    for n in negocios_linhas:
        nid = n.get("ID") or ""
        if not nid:
            continue
        info = lead_info.get(n.get("ID Lead") or "", {})
        local = {k: info.get(k) for k in ("praca", "supervisao_id", "unidade_id")} if info else onde(n, None, None)
        existente = leads_db.get(nid)
        if existente:
            cont["ja_importados"] += 1
            mudar("leads", existente, {**local, "pixel": info.get("pixel")})
            continue
        fp = funis_planilha.get(n.get("Funil") or "")
        if not fp:
            avisos.append(f"Negócio {nid}: funil {n.get('Funil')} não encontrado na aba FUNIS.")
            continue
        etapas = fp["funil"]["etapas"]
        etapa = fp["etapas"].get(n.get("Etapa")) or etapas[0]
        pos = next((i for i, e in enumerate(etapas) if e["id"] == etapa["id"]), 0)
        status = (n.get("Status") or "aberto").strip().lower()
        status = status if status in ("aberto", "ganho", "perdido") else "aberto"
        tags = [info.get("tipo") or ("Evento" if "EVENTO" in (n.get("Funil") or "") else "Revenda")]
        loja = n.get("Loja") or ""
        if loja:
            tags.append(f"Loja {_loja_curta(loja)}")
        tags += [t for t in n.get("_extra", []) if len(t) <= 20]  # temperatura (Frio, Morno, Quente)
        etiquetas.update(tags)
        criado = _data(n.get("Criado em")) or agora
        doc = {
            "id": negocio_novo_id[nid], "nome": (info.get("nome") or n.get("ID Lead") or nid)[:1000], "cliente_id": cliente_do_lead.get(n.get("ID Lead")),
            "corretor_id": pessoa_do_usuario.get(n.get("Responsável ID") or ""), "origem": info.get("origem") or "Planilha",
            "estagio": status if status != "aberto" else "novo", "valor_estimado": _valor(n.get("Valor")),
            "observacoes": None, "funil_id": fp["funil"]["id"], "etapa_id": etapa["id"], "status": status,
            "etapa_desde": _data(n.get("Na etapa desde")) or criado,
            "motivo_perda": ((n.get("Motivo da perda") or "").strip()[:200] or None) if status == "perdido" else None,
            "previsao_fechamento": None, "etiquetas": tags, "entrada_id": (entradas_db.get(n.get("ID Lead")) or {}).get("id")
            or next((e["id"] for e in novas_entradas if e["planilha_id"] == n.get("ID Lead")), None),
            "etapas_alcancadas": [e["id"] for e in etapas[:pos + 1]], "pixel": info.get("pixel"), **local,
            "planilha_id": nid, "created_at": criado, "updated_at": _data(n.get("Atualizado em")) or criado,
            "closed_at": _data(n.get("Ganho em") if status == "ganho" else n.get("Perdido em")) if status != "aberto" else None,
        }
        novos_negocios.append(doc)
        novos_hist.append({"id": new_id(), "negocio_id": doc["id"], "tipo": "criado", "texto": "Negócio importado da planilha CRM Revendas",
                           "autor": por, "autor_id": None, "fixado": False, "em": criado})
        cont["negocios"] += 1
    if not simular and novos_negocios:
        await banco.leads.insert_many(novos_negocios)
    negocio_id = {nid: d["id"] for nid, d in leads_db.items()} | {d["planilha_id"]: d["id"] for d in novos_negocios}
    negocio_dono = {d["id"]: (d.get("corretor_id"), d.get("cliente_id")) for d in [*leads_db.values(), *novos_negocios]}

    # ------------------------------------------------------------ atividades
    novas_ativ = []
    nomes_drive = {(l.get("ID arquivo") or "").strip(): (l.get("Nome do arquivo") or "").strip() for l in _tabela(abas, "ANEXOS")}
    pendentes: list[dict] = []

    def prints_da(linha: dict, ativ: dict) -> None:
        for fid in re.split(r"[,;\s]+", linha.get("Anexo (ID arquivo)") or ""):
            if not fid or fid in drive_ja:
                continue
            drive_ja.add(fid)
            cont["prints_pendentes"] += 1
            pendentes.append({"id": new_id(), "drive_id": fid, "nome": nomes_drive.get(fid) or None, "atividade_id": ativ["id"],
                              "atividade_assunto": ativ.get("assunto"), "negocio_id": ativ.get("negocio_id"),
                              "autor_pessoa_id": ativ.get("corretor_id"), "autor_nome": linha.get("Feito por") or None, "created_at": agora})

    for a in _tabela(abas, "ATIVIDADES"):
        aid = a.get("ID Atividade") or ""
        if not aid:
            continue
        if aid in ativ_db:
            cont["ja_importados"] += 1
            prints_da(a, ativ_db[aid])
            continue
        neg = negocio_id.get(a.get("ID Negócio") or "")
        dono_neg, cliente_neg = negocio_dono.get(neg, (None, None))
        situacao = (a.get("Status") or "").strip()
        feita = bool(a.get("Concluída em")) or _sem_acento(situacao).lower() in {_sem_acento(x).lower() for x in FEITAS}
        notas = [a.get("Observação") or "", f"Resultado: {situacao}" if situacao else "",
                 f"Anexo no Google Drive da planilha (ID {a.get('Anexo (ID arquivo)')})" if a.get("Anexo (ID arquivo)") else ""]
        tipo = TIPOS_ATIVIDADE.get(_sem_acento(a.get("Tipo") or "").lower(), "tarefa")
        novas_ativ.append({
            "id": new_id(), "tipo": tipo, "assunto": (a.get("Título") or a.get("Tipo") or "Atividade")[:300],
            "data": _dia(a.get("Prazo")) or _dia(a.get("Criada em")) or agora.date().isoformat(), "hora": None, "duracao_min": 30,
            "negocio_id": neg, "entrada_id": None, "pessoa_id": cliente_neg, "imovel_id": None,
            "corretor_id": pessoa_por_nome.get(_sem_acento(a.get("Feito por") or "")) or dono_neg,
            "notas": "\n".join(x for x in notas if x)[:5000] or None, "concluida": feita,
            "resultado": "positivo" if _sem_acento(situacao).lower() in ("sucesso", "realizada") else ("negativo" if situacao and feita is False else None),
            "concluida_em": _data(a.get("Concluída em")) if feita else None, "created_by": a.get("Criada por") or "Planilha",
            "planilha_id": aid, "created_at": _data(a.get("Criada em")) or agora, "updated_at": agora,
        })
        cont["atividades"] += 1
        prints_da(a, novas_ativ[-1])
    if not simular and novas_ativ:
        await banco.atividades.insert_many(novas_ativ)
    if not simular and pendentes:
        await banco.anexos_pendentes.insert_many(pendentes)

    # ------------------------------------------------------------ histórico
    for i, h in enumerate(_tabela(abas, "HISTORICO")):
        neg = negocio_id.get(h.get("ID Negócio") or "")
        if not neg:
            continue
        chave = f"{h.get('ID Negócio')}|{h.get('Data')}|{h.get('Ação')}|{h.get('Detalhe')}|{i}"
        if chave in hist_db:
            continue
        acao = h.get("Ação") or ""
        a = _sem_acento(acao).lower()
        tipo = "status" if ("perdido" in a or "ganho" in a) else "etapa" if "etapa" in a else "atividade" if "atividade" in a else "campo"
        cont["historico"] += 1
        novos_hist.append({"id": new_id(), "negocio_id": neg, "tipo": tipo, "texto": f"{acao}: {h.get('Detalhe')}" if h.get("Detalhe") else acao,
                           "autor": h.get("Usuário") or None, "autor_id": None, "fixado": False, "em": _data(h.get("Data")) or agora, "planilha_id": chave})
    if not simular and novos_hist:
        await banco.historico.insert_many(novos_hist)

    # ------------------------------------------------------------ origens, etiquetas e o que só faltava completar
    if not simular and (origens or etiquetas):
        crm = cfg.get("crm") or {}
        novas_origens = list(dict.fromkeys([*(crm.get("origens") or []), *sorted(origens)]))[:60]
        novas_etiquetas = list(dict.fromkeys([*(crm.get("etiquetas") or []), *sorted(etiquetas)]))[:80]
        await banco.configuracoes.update_one({"id": "singleton"}, {"$set": {"crm.origens": novas_origens, "crm.etiquetas": novas_etiquetas}})
    if not simular and not leads_db and not entradas_db:
        # no CRM da planilha, contato e reunião só contavam com o print anexado (só na 1ª importação: depois vale o que a empresa escolher)
        await banco.configuracoes.update_one({"id": "singleton"}, {"$set": {"crm.print_obrigatorio": True}})
    if atualizar:
        await _em_paralelo(atualizar)

    # ------------------------------------------------------------ pedidos e comissões (mantém o que a planilha decidiu)
    ped_db = {d["planilha_id"] for d in await banco.pedidos.find({"planilha_id": {"$ne": None}}, {"planilha_id": 1}).to_list(50000) if d.get("planilha_id")}
    status_ped = {"PENDENTE": "pendente", "APROVADA": "aprovada", "PAGA": "paga", "CANCELADA": "cancelado", "CANCELADO": "cancelado"}
    novos_ped = []
    neg_por_id = {d["id"]: d for d in [*leads_db.values(), *novos_negocios]}
    for p in _tabela(abas, "PEDIDOS"):
        pid = p.get("ID") or ""
        neg = neg_por_id.get(negocio_id.get(p.get("ID Negócio") or "") or "")
        if not pid or not neg:
            continue
        if pid in ped_db:
            cont["ja_importados"] += 1
            continue
        dia = _dia(p.get("Data do pedido")) or agora.date().isoformat()
        st = status_ped.get(_sem_acento(p.get("Status comissão") or ""), "pendente")
        pct = _valor((p.get("% Comissão") or "").replace("%", "")) or 0.0
        novos_ped.append({
            "id": new_id(), "negocio_id": neg["id"], "negocio_nome": neg.get("nome"), "cliente_id": neg.get("cliente_id"),
            "cliente_chave": neg.get("cliente_id") or f"negocio:{neg['id']}", "numero": (p.get("Nº do pedido") or "").strip() or pid,
            "data": dia, "mes": dia[:7], "valor": _valor(p.get("Valor")) or 0.0,
            "beneficiario_id": pessoa_do_usuario.get(p.get("Beneficiário ID") or "") or neg.get("corretor_id"),
            "praca": neg.get("praca"), "supervisao_id": neg.get("supervisao_id"), "unidade_id": neg.get("unidade_id"),
            "status": st if pct or st == "cancelado" else "sem_comissao", "ordem": int(_valor(p.get("Ordem do cliente")) or 0) or None,
            "pct": pct, "comissao": _valor(p.get("Comissão")) or 0.0, "observacao": (p.get("Observação") or "").strip() or None,
            "aprovado_por": p.get("Aprovado por") or None, "aprovado_em": _data(p.get("Aprovado em")),
            "pago_por": p.get("Pago por") or None, "pago_em": _data(p.get("Pago em")),
            "registrado_por": p.get("Registrado por") or "Planilha", "registrado_em": _data(p.get("Registrado em")) or agora,
            "planilha_id": pid,
        })
        cont["pedidos"] = cont.get("pedidos", 0) + 1
    if novos_ped and not simular:
        await banco.pedidos.insert_many(novos_ped)
    if cont["prints_pendentes"]:
        avisos.append(f"{cont['prints_pendentes']} print(s) do Google Drive estão ligados às atividades. Baixe a pasta de anexos do Drive "
                      "e envie os arquivos em Empresas, Importar prints.")
    if any(a["link"] for a in acessos):
        avisos.append("Os links de acesso valem por 24 horas e são de uso único. Mande para cada pessoa hoje.")
    return {"simulado": simular, "contagens": cont, "acessos": acessos, "avisos": avisos}
