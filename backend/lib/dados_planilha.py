"""Planilha de dados da empresa: modelo para preencher, exportação e importação (todos os segmentos).

Abas: LEIA-ME, CLIENTES (PACIENTES em clínica), LEADS e NEGOCIOS (só nos segmentos com CRM) e LISTAS
(escondida, alimenta as listas de escolha). A planilha é personalizada pela empresa: funis, etapas,
origens e e-mails da equipe dela viram listas de escolha, e os nomes seguem o segmento.

Importar não duplica: linha com "ID SAX" (vem na exportação) atualiza o registro; sem ID, o cliente é
procurado por CPF/CNPJ, telefone e e-mail, e o lead pelo telefone. Linha com problema volta em `erros`
com a aba e o número da linha, sem travar as outras. Lê o que existe uma vez e grava em lote, para caber
nos 60 s da função no Vercel.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime, timezone

from lib.segmentos import SEGMENTOS, SEGMENTO_PADRAO
from models.common import new_id, now_utc

MAX_LINHAS = 5000
COL_ID = "ID SAX"
MAX_MENSAGENS = 200


class ErroPlanilha(Exception):
    pass


# ------------------------------------------------------------------ colunas por segmento


def _seg(segmento: str | None) -> dict:
    return SEGMENTOS.get(segmento or "") or SEGMENTOS[SEGMENTO_PADRAO]


def tem_crm(segmento: str | None) -> bool:
    return "crm" in _seg(segmento)["modulos"]


def _item(segmento: str | None) -> str | None:
    """Produto do negócio que tem código no sistema (imóvel ou veículo)."""
    return {"imobiliaria": "imovel", "veiculos": "veiculo"}.get(segmento or "")


def nome_aba_clientes(segmento: str | None) -> str:
    return _sem_acento(_seg(segmento)["termos"]["clientes"]).upper()


def colunas(segmento: str | None) -> dict[str, list[tuple[str, str]]]:
    """aba -> [(chave, cabeçalho)]. O asterisco marca o campo obrigatório."""
    t = _seg(segmento)["termos"]
    item = _item(segmento)
    cli = [("id", COL_ID), ("nome", "Nome*"), ("telefone", "Telefone"), ("telefone2", "Telefone 2"), ("email", "E-mail"),
           ("cpf_cnpj", "CPF ou CNPJ"), ("data_nascimento", "Data de nascimento"), ("profissao", "Profissão"),
           ("cep", "CEP"), ("logradouro", "Endereço"), ("numero", "Número"), ("complemento", "Complemento"),
           ("bairro", "Bairro"), ("cidade", "Cidade"), ("estado", "UF"), ("observacoes", "Observações")]
    if segmento == "imobiliaria":
        cli.insert(2, ("papel", "Tipo"))
    if segmento == "comercio":
        cli.insert(2, ("nome_fantasia", "Estabelecimento"))
    abas = {nome_aba_clientes(segmento): cli}
    if tem_crm(segmento):
        leads = [("id", COL_ID), ("nome", "Nome*"), ("telefone", "Telefone"), ("email", "E-mail"), ("origem", "Origem"),
                 ("temperatura", "Temperatura")]
        if segmento == "imobiliaria":
            leads.append(("interesse", "Interesse"))
        if item:
            leads.append(("item", f"Código do {t['item'].lower()}"))
        leads += [("valor", "Valor estimado"), ("responsavel", "Responsável (e-mail)"), ("situacao", "Situação"),
                  ("mensagem", "Mensagem"), ("criado", "Data de entrada")]
        neg = [("id", COL_ID), ("nome", "Título*"), ("cliente", t["cliente"]), ("cliente_telefone", f"Telefone do {t['cliente'].lower()}"),
               ("cliente_email", f"E-mail do {t['cliente'].lower()}"), ("funil", "Funil"), ("etapa", "Etapa"), ("status", "Situação"),
               ("valor", "Valor"), ("origem", "Origem"), ("temperatura", "Temperatura"), ("responsavel", "Responsável (e-mail)"),
               ("previsao", "Previsão de fechamento")]
        if item:
            neg.append(("item", f"Código do {t['item'].lower()}"))
        neg += [("etiquetas", "Etiquetas (separe por vírgula)"), ("motivo_perda", "Motivo da perda"),
                ("observacoes", "Observações"), ("criado", "Criado em")]
        abas["LEADS"] = leads
        abas["NEGOCIOS"] = neg
    return abas


STATUS_NEG = {"em andamento": "aberto", "aberto": "aberto", "andamento": "aberto", "ganho": "ganho", "ganha": "ganho",
              "vendido": "ganho", "fechado": "ganho", "perdido": "perdido", "perdida": "perdido"}
STATUS_NEG_ROTULO = {"aberto": "Em andamento", "ganho": "Ganho", "perdido": "Perdido"}
SITUACAO_LEAD = {"novo": "novo", "em contato": "em_contato", "em atendimento": "em_contato", "descartado": "descartado",
                 "convertido": "convertido"}
SITUACAO_LEAD_ROTULO = {"novo": "Novo", "em_contato": "Em contato", "descartado": "Descartado", "convertido": "Convertido em negócio"}
INTERESSE = {"compra": "compra", "comprar": "compra", "locacao": "locacao", "alugar": "locacao", "aluguel": "locacao",
             "venda": "venda", "vender": "venda", "outro": "outro"}
INTERESSE_ROTULO = {"compra": "Compra", "locacao": "Locação", "venda": "Venda", "outro": "Outro"}
PAPEL_IMOB = {"cliente": ["cliente"], "proprietario": ["proprietario"], "cliente e proprietario": ["cliente", "proprietario"]}


# ------------------------------------------------------------------ texto


def _sem_acento(t: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", t or "") if unicodedata.category(c) != "Mn")


def _chave(t: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", " ", _sem_acento(t or "").lower()).strip()


def _txt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _digitos(v: str | None) -> str:
    return re.sub(r"\D", "", v or "")


def _dia(v: str | None) -> str | None:
    """DD/MM/AAAA, DD/MM/AA ou AAAA-MM-DD -> AAAA-MM-DD. Texto ruim -> None."""
    v = (v or "").strip()
    if not v:
        return None
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", v)
    if m:
        a, mes, d = (int(x) for x in m.groups())
    else:
        m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", v)
        if not m:
            return None
        d, mes, a = (int(x) for x in m.groups())
        if a < 100:
            a += 2000 if a <= date.today().year % 100 else 1900
    try:
        return date(a, mes, d).isoformat()
    except ValueError:
        return None


def _valor(v: str | None) -> float | None:
    v = (v or "").replace("R$", "").replace(" ", "").strip()
    if not v:
        return None
    if "," in v:
        v = v.replace(".", "").replace(",", ".")
    try:
        n = round(float(v), 2)
    except ValueError:
        return None
    return n if n >= 0 else None


def _br(iso: str | None) -> str:
    if not iso:
        return ""
    try:
        return date.fromisoformat(str(iso)[:10]).strftime("%d/%m/%Y")
    except ValueError:
        return ""


def _br_dt(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    return _br(v if isinstance(v, str) else None)


# ------------------------------------------------------------------ leitura do arquivo


def ler(conteudo: bytes) -> dict[str, list[list[str]]]:
    """Abas -> linhas de texto. Aceita .xlsx (Excel ou Google Planilhas) e .csv (uma aba só)."""
    if conteudo[:2] == b"PK":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
        except Exception:
            raise ErroPlanilha("Não consegui abrir o arquivo. Salve como Excel (.xlsx) e envie de novo.")
        abas = {}
        for ws in wb.worksheets:
            linhas = []
            for linha in ws.iter_rows(values_only=True):
                linhas.append([_txt(c) for c in linha])
                if len(linhas) > MAX_LINHAS + 1:
                    raise ErroPlanilha(f"A aba {ws.title} passa de {MAX_LINHAS} linhas. Divida em arquivos menores.")
            abas[_sem_acento(ws.title).strip().upper()] = linhas
        wb.close()
        return abas
    for cod in ("utf-8-sig", "latin-1"):
        try:
            texto = conteudo.decode(cod)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ErroPlanilha("Formato não reconhecido. Envie .xlsx ou .csv.")
    if "\x00" in texto[:2000]:
        raise ErroPlanilha("Formato não reconhecido. Envie .xlsx ou .csv.")
    amostra = texto[:4000]
    sep = ";" if amostra.count(";") > amostra.count(",") else ","
    linhas = [[c.strip() for c in l] for l in csv.reader(io.StringIO(texto), delimiter=sep)]
    if len(linhas) > MAX_LINHAS + 1:
        raise ErroPlanilha(f"O arquivo passa de {MAX_LINHAS} linhas. Divida em arquivos menores.")
    return {"CSV": linhas}


def _qual_aba(nome: str, cab: list[str], segmento: str | None) -> str | None:
    n = _chave(nome)
    if n in ("clientes", "pacientes", "contatos", "pessoas", _chave(nome_aba_clientes(segmento))):
        return "clientes"
    if n in ("leads", "lead", "entradas"):
        return "leads"
    if n in ("negocios", "negocio", "oportunidades", "vendas"):
        return "negocios"
    if n == "csv":
        h = {_chave(c) for c in cab}
        if {"titulo", "funil", "etapa"} & h:
            return "negocios"
        if "origem" in h or "temperatura" in h:
            return "leads"
        return "clientes"
    return None


def _mapear(cab: list[str], cols: list[tuple[str, str]]) -> dict[int, str]:
    """posição da coluna -> chave. Aceita o cabeçalho do modelo, a chave ou variações sem acento."""
    gerais = {"id": "id", "id sax": "id", "celular": "telefone", "whatsapp": "telefone", "fone": "telefone",
                 "email": "email", "cpf": "cpf_cnpj", "cnpj": "cpf_cnpj", "documento": "cpf_cnpj", "nascimento": "data_nascimento",
                 "uf": "estado", "estado": "estado", "rua": "logradouro", "endereco": "logradouro", "obs": "observacoes",
                 "titulo": "nome", "negocio": "nome", "responsavel": "responsavel", "vendedor": "responsavel", "corretor": "responsavel",
                 "status": "status", "situacao": "status", "valor": "valor", "previsao": "previsao", "etiquetas": "etiquetas",
                 "data": "criado", "data de entrada": "criado", "criado em": "criado"}
    apelidos = gerais | {_chave(chave): chave for chave, _ in cols} | {_chave(titulo): chave for chave, titulo in cols}
    saida = {}
    for i, c in enumerate(cab):
        k = apelidos.get(_chave(c.replace("*", "")))
        if k and k in {ch for ch, _ in cols} and k not in saida.values():
            saida[i] = k
    return saida


# ------------------------------------------------------------------ geração do .xlsx


async def _contexto(banco) -> dict:
    """O que personaliza a planilha: funis e etapas, origens, equipe e códigos dos itens."""
    funis = await banco.funis.find({}).sort("ordem", 1).to_list(100)
    cfg = (await banco.configuracoes.find_one({"id": "singleton"}) or {})
    crm = cfg.get("crm") or {}
    usuarios = [u for u in await banco.usuarios.find({}, {"email": 1, "nome": 1, "pessoa_id": 1, "ativo": 1}).to_list(2000) if u.get("email")]
    return {"funis": funis, "origens": crm.get("origens") or [], "motivos": crm.get("motivos_perda") or [],
            "etiquetas": crm.get("etiquetas") or [], "usuarios": usuarios, "empresa": cfg.get("nome_software") or "",
            "regras_temperatura": crm.get("regras_temperatura") or []}


def _xlsx(segmento: str | None, ctx: dict, dados: dict[str, list[list]] | None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    seg = _seg(segmento)
    t = seg["termos"]
    cols = colunas(segmento)
    wb = Workbook()
    leia = wb.active
    leia.title = "LEIA-ME"
    roxo, laranja = "4A03A2", "FF7A00"
    aba_cli = nome_aba_clientes(segmento)

    # listas de escolha (aba escondida, criada no fim)
    etapas = sorted({e["nome"] for f in ctx["funis"] for e in f.get("etapas", [])})
    valores = {
        "funil": [f["nome"] for f in ctx["funis"]], "etapa": etapas, "origem": ctx["origens"],
        "temperatura": ["Quente", "Morno", "Frio"], "status": list(STATUS_NEG_ROTULO.values()),
        "situacao": ["Novo", "Em contato", "Descartado"], "responsavel": [u["email"] for u in ctx["usuarios"] if u.get("ativo", True)],
        "interesse": list(INTERESSE_ROTULO.values()), "papel": ["Cliente", "Proprietário", "Cliente e proprietário"],
        "motivo_perda": ctx["motivos"], "estado": ["AC", "AL", "AP", "AM", "BA", "CE", "DF", "ES", "GO", "MA", "MT", "MS", "MG", "PA",
                                                  "PB", "PR", "PE", "PI", "RJ", "RN", "RS", "RO", "RR", "SC", "SP", "SE", "TO"],
    }
    faixa: dict[str, str] = {}
    usadas = [(k, v) for k, v in valores.items() if v]
    for j, (chave, itens) in enumerate(usadas):
        letra = get_column_letter(j + 1)
        faixa[chave] = f"LISTAS!${letra}$2:${letra}${len(itens) + 1}"

    larguras = {"id": 14, "nome": 32, "observacoes": 40, "mensagem": 40, "logradouro": 30, "email": 28, "responsavel": 28,
                "cliente": 28, "cliente_email": 28, "etiquetas": 28, "funil": 18, "etapa": 24}
    exemplos = {"telefone": "(41) 99999-0000", "data_nascimento": "DD/MM/AAAA", "criado": "DD/MM/AAAA", "previsao": "DD/MM/AAAA",
                "valor": "1500,00", "cpf_cnpj": "Só números ou com pontos"}
    for nome_aba, cs in cols.items():
        ws = wb.create_sheet(nome_aba)
        for j, (chave, titulo) in enumerate(cs, start=1):
            cel = ws.cell(row=1, column=j, value=titulo)
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="9B9B9B" if chave == "id" else roxo)
            cel.alignment = Alignment(vertical="center")
            if chave == "id":
                cel.comment = Comment("Preenchido pelo sistema na exportação. Deixe vazio em linha nova; não mude o que veio.", "SAX")
            elif chave in exemplos:
                cel.comment = Comment(f"Ex.: {exemplos[chave]}", "SAX")
            ws.column_dimensions[get_column_letter(j)].width = larguras.get(chave, 18)
            lista = faixa.get({"status": "status" if nome_aba == "NEGOCIOS" else "situacao", "situacao": "situacao"}.get(chave, chave))
            if lista:
                dv = DataValidation(type="list", formula1=lista, allow_blank=True, showErrorMessage=chave not in ("origem", "etapa"))
                dv.error = "Escolha um valor da lista."
                ws.add_data_validation(dv)
                dv.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{MAX_LINHAS + 1}")
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 22
        for linha in (dados or {}).get(nome_aba, []):
            ws.append(linha)

    linhas_leia = [
        (f"Dados de {ctx['empresa'] or 'sua empresa'} no SAX CRM", True),
        (f"Segmento: {seg['nome']}. Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}.", False),
        ("", False),
        ("COMO PREENCHER", True),
        (f"1. Uma linha por registro nas abas {', '.join(cols)}. Só o campo com * é obrigatório.", False),
        ("2. Não mude os nomes das abas nem a primeira linha (cabeçalhos).", False),
        ("3. A coluna ID SAX é do sistema: vazia em linha nova; na planilha exportada, não mude. Com ID, a linha atualiza o registro.", False),
        (f"4. Sem ID, o {t['cliente'].lower()} é procurado por CPF/CNPJ, telefone e e-mail e o lead pelo telefone: importar de novo não duplica.", False),
        ("5. Datas no formato DD/MM/AAAA. Valores como 1500,00. Responsável é o e-mail de login da pessoa da equipe.", False),
        ("6. Em Importar dados, use Conferir antes: o sistema mostra o que vai entrar e as linhas com problema, sem gravar nada.", False),
    ]
    if tem_crm(segmento):
        linhas_leia += [
            ("", False),
            ("LEADS E NEGÓCIOS", True),
            ("Temperatura: Quente, Morno ou Frio. Em branco, a origem decide pelas regras do CRM.", False),
            ("Funil e etapa: use os nomes abaixo. Funil em branco vai para o principal; etapa em branco, para a primeira.", False),
            ("Situação do negócio: Em andamento, Ganho ou Perdido.", False),
            ("", False),
            ("FUNIS E ETAPAS DESTA EMPRESA", True),
            *[(f"{f['nome']}: {', '.join(e['nome'] for e in f.get('etapas', []))}", False) for f in ctx["funis"]],
            ("", False),
            ("ORIGENS", True),
            (", ".join(ctx["origens"]) or "Livre", False),
        ]
    linhas_leia += [("", False), ("EQUIPE (e-mail para a coluna Responsável)", True),
                    *[(f"{u.get('nome') or ''}: {u['email']}", False) for u in ctx["usuarios"] if u.get("ativo", True)]]
    for i, (texto, titulo) in enumerate(linhas_leia, start=1):
        c = leia.cell(row=i, column=1, value=texto)
        if titulo:
            c.font = Font(bold=True, color=laranja if i > 1 else roxo, size=14 if i == 1 else 11)
    leia.column_dimensions["A"].width = 120
    listas = wb.create_sheet("LISTAS")
    for j, (chave, itens) in enumerate(usadas):
        letra = get_column_letter(j + 1)
        listas[f"{letra}1"] = chave
        for i, v in enumerate(itens, start=2):
            listas[f"{letra}{i}"] = v
    listas.sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def modelo(banco, segmento: str | None) -> bytes:
    return _xlsx(segmento, await _contexto(banco), None)


async def exportar(banco, segmento: str | None) -> bytes:
    ctx = await _contexto(banco)
    cols = colunas(segmento)
    aba_cli = nome_aba_clientes(segmento)
    email_da_pessoa = {u["pessoa_id"]: u["email"] for u in ctx["usuarios"] if u.get("pessoa_id")}
    equipe = set(email_da_pessoa)
    pessoas = await banco.pessoas.find({}).to_list(None)
    por_id = {p["id"]: p for p in pessoas}
    item = _item(segmento)
    codigos = {}
    if item:
        colecao = banco.imoveis if item == "imovel" else banco.veiculos
        codigos = {d["id"]: d.get("codigo") or "" for d in await colecao.find({}, {"id": 1, "codigo": 1}).to_list(None)}
    funis = {f["id"]: f for f in ctx["funis"]}
    dados: dict[str, list[list]] = {}

    def linha(cs, valores: dict) -> list:
        return [valores.get(k, "") if valores.get(k) is not None else "" for k, _ in cs]

    clientes = []
    for p in pessoas:
        papeis = set(p.get("papeis") or [])
        if p["id"] in equipe and not (papeis & {"cliente", "paciente", "proprietario"}):
            continue
        if papeis and not (papeis & {"cliente", "paciente", "proprietario"}):
            continue
        v = {k: p.get(k) for k, _ in cols[aba_cli]}
        v["id"] = p["id"]
        v["data_nascimento"] = _br(p.get("data_nascimento"))
        if segmento == "imobiliaria":
            v["papel"] = ("Cliente e proprietário" if {"cliente", "proprietario"} <= papeis else "Proprietário" if "proprietario" in papeis else "Cliente")
        clientes.append(linha(cols[aba_cli], v))
    dados[aba_cli] = clientes
    if tem_crm(segmento):
        leads = []
        for e in await banco.entradas.find({}).sort("created_at", 1).to_list(None):
            leads.append(linha(cols["LEADS"], {
                "id": e["id"], "nome": e.get("nome"), "telefone": e.get("telefone"), "email": e.get("email"), "origem": e.get("origem"),
                "temperatura": (e.get("temperatura") or "").capitalize(), "interesse": INTERESSE_ROTULO.get(e.get("interesse") or "", ""),
                "item": codigos.get(e.get("imovel_id") or e.get("veiculo_id") or "", ""), "valor": e.get("valor_estimado"),
                "responsavel": email_da_pessoa.get(e.get("corretor_id") or "", ""), "situacao": SITUACAO_LEAD_ROTULO.get(e.get("status") or "", ""),
                "mensagem": e.get("mensagem"), "criado": _br_dt(e.get("created_at")),
            }))
        dados["LEADS"] = leads
        negs = []
        for n in await banco.leads.find({}).sort("created_at", 1).to_list(None):
            f = funis.get(n.get("funil_id") or "") or {}
            etapa = next((e["nome"] for e in f.get("etapas", []) if e["id"] == n.get("etapa_id")), "")
            cli = por_id.get(n.get("cliente_id") or "") or {}
            negs.append(linha(cols["NEGOCIOS"], {
                "id": n["id"], "nome": n.get("nome"), "cliente": cli.get("nome"), "cliente_telefone": cli.get("telefone"),
                "cliente_email": cli.get("email"), "funil": f.get("nome"), "etapa": etapa,
                "status": STATUS_NEG_ROTULO.get(n.get("status") or "aberto", "Em andamento"), "valor": n.get("valor_estimado"),
                "origem": n.get("origem"), "temperatura": (n.get("temperatura") or "").capitalize(),
                "responsavel": email_da_pessoa.get(n.get("corretor_id") or "", ""), "previsao": _br(n.get("previsao_fechamento")),
                "item": codigos.get(n.get("imovel_id") or n.get("veiculo_id") or "", ""), "etiquetas": ", ".join(n.get("etiquetas") or []),
                "motivo_perda": n.get("motivo_perda"), "observacoes": n.get("observacoes"), "criado": _br_dt(n.get("created_at")),
            }))
        dados["NEGOCIOS"] = negs
    return _xlsx(segmento, ctx, dados)


# ------------------------------------------------------------------ importação


async def importar(banco, segmento: str | None, conteudo: bytes, *, simular: bool, por: str, por_usuario_id: str | None = None) -> dict:
    from lib.crm import estagio_legado
    from lib.temperatura import normalizar as normalizar_temp
    from lib.temperatura import pela_origem
    from models.pessoas import Pessoa, cnpj_valido, cpf_valido, formatar_documento

    abas = ler(conteudo)
    cols = colunas(segmento)
    aba_cli = nome_aba_clientes(segmento)
    tabelas: dict[str, list[tuple[int, dict]]] = {}
    for nome, linhas in abas.items():
        cheias = [(i, l) for i, l in enumerate(linhas, start=1) if any(c for c in l)]
        if not cheias:
            continue
        _, cab = cheias[0]
        tipo = _qual_aba(nome, cab, segmento)
        if not tipo or (tipo != "clientes" and not tem_crm(segmento)):
            continue
        chave_cols = cols[aba_cli] if tipo == "clientes" else cols["LEADS" if tipo == "leads" else "NEGOCIOS"]
        mapa = _mapear(cab, chave_cols)
        if "nome" not in mapa.values():
            raise ErroPlanilha(f"A aba {nome} não tem a coluna Nome (ou Título). Use o modelo baixado do sistema.")
        tabelas[tipo] = [(n, {mapa[j]: (l[j] if j < len(l) else "").strip() for j in mapa}) for n, l in cheias[1:]]
    if not tabelas:
        raise ErroPlanilha("Não encontrei as abas do modelo (" + ", ".join(cols) + "). Baixe o modelo e preencha nele.")

    agora = now_utc()
    cont = {k: 0 for k in ("clientes_novos", "clientes_atualizados", "leads_novos", "leads_atualizados",
                           "negocios_novos", "negocios_atualizados", "linhas_com_erro")}
    erros: list[dict] = []
    avisos: list[dict] = []

    def erro(aba: str, linha: int, motivo: str) -> None:
        cont["linhas_com_erro"] += 1
        if len(erros) < MAX_MENSAGENS:
            erros.append({"aba": aba, "linha": linha, "motivo": motivo})

    def aviso(aba: str, linha: int, motivo: str) -> None:
        if len(avisos) < MAX_MENSAGENS:
            avisos.append({"aba": aba, "linha": linha, "motivo": motivo})

    # ---- o que já existe (uma leitura por coleção)
    ctx = await _contexto(banco)
    pessoas = await banco.pessoas.find({}, {"id": 1, "nome": 1, "telefone_digitos": 1, "telefone": 1, "documento_digitos": 1,
                                            "email": 1, "papeis": 1}).to_list(None)
    p_id = {p["id"]: p for p in pessoas}
    p_doc = {p["documento_digitos"]: p for p in pessoas if p.get("documento_digitos")}
    p_tel: dict[str, dict] = {}
    p_mail: dict[str, dict] = {}
    p_nome: dict[str, dict] = {}

    def indexar(p: dict) -> None:
        tel = p.get("telefone_digitos") or _digitos(p.get("telefone"))
        if len(tel) >= 10:
            p_tel.setdefault(tel[-11:], p)
        if p.get("email"):
            p_mail.setdefault(p["email"].lower(), p)
        p_nome.setdefault(_chave(p.get("nome")), p)
        p_id[p["id"]] = p
        if p.get("documento_digitos"):
            p_doc[p["documento_digitos"]] = p

    for p in pessoas:
        indexar(p)
    usuario_por_email = {u["email"].lower(): u for u in ctx["usuarios"]}
    usuario_por_nome = {_chave(u.get("nome")): u for u in ctx["usuarios"] if u.get("nome")}
    funis = ctx["funis"]
    funil_padrao = next((f for f in funis if f.get("padrao")), funis[0] if funis else None)
    regras = ctx["regras_temperatura"]
    papel_cli = "paciente" if segmento in ("terapia", "odontologia") else "cliente"
    item = _item(segmento)
    cod_item: dict[str, str] = {}
    if item:
        colecao = banco.imoveis if item == "imovel" else banco.veiculos
        cod_item = {_chave(d.get("codigo")): d["id"] for d in await colecao.find({}, {"id": 1, "codigo": 1}).to_list(None) if d.get("codigo")}

    novas_pessoas: list[dict] = []
    novos_ids: set[str] = set()
    atualizar: list = []

    def responsavel(aba: str, n: int, v: str) -> str | None:
        if not v:
            return None
        u = usuario_por_email.get(v.lower()) or usuario_por_nome.get(_chave(v))
        if not u or not u.get("pessoa_id"):
            aviso(aba, n, f"Responsável \"{v}\" não está na equipe: ficou sem responsável.")
            return None
        return u["pessoa_id"]

    def achar_pessoa(doc: str = "", tel: str = "", email: str = "", nome: str = "") -> dict | None:
        return (p_doc.get(doc) if doc else None) or (p_tel.get(tel[-11:]) if len(tel) >= 10 else None) \
            or (p_mail.get(email.lower()) if email else None) or (p_nome.get(_chave(nome)) if nome and not (tel or email or doc) else None)

    def nova_pessoa(dados: dict) -> dict:
        p = Pessoa(**dados).model_dump()
        p["telefone_digitos"] = _digitos(p.get("telefone")) or None
        p["documento_digitos"] = _digitos(p.get("cpf_cnpj")) or None
        p["criado_por_id"] = por_usuario_id
        p["created_by"] = por_usuario_id
        p["origem_importacao"] = True
        novas_pessoas.append(p)
        novos_ids.add(p["id"])
        indexar(p)
        return p

    # ---- clientes / pacientes
    for n, l in tabelas.get("clientes", []):
        nome = l.get("nome", "")
        if not nome:
            erro(aba_cli, n, "Falta o nome.")
            continue
        dados = {k: l.get(k) or None for k in ("telefone", "telefone2", "profissao", "cep", "logradouro", "numero", "complemento",
                                                "bairro", "cidade", "observacoes", "nome_fantasia")}
        dados["nome"] = nome[:1000]
        email = (l.get("email") or "").strip().lower()
        if email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
            aviso(aba_cli, n, f"E-mail \"{email}\" inválido: ficou em branco.")
            email = ""
        dados["email"] = email or None
        doc = _digitos(l.get("cpf_cnpj"))
        if doc and not ((len(doc) == 11 and cpf_valido(doc)) or (len(doc) == 14 and cnpj_valido(doc))):
            aviso(aba_cli, n, "CPF/CNPJ inválido: ficou em branco.")
            doc = ""
        if doc:
            dados["cpf_cnpj"] = formatar_documento(doc)
            dados["tipo_pessoa"] = "pj" if len(doc) == 14 else "pf"
        nasc = l.get("data_nascimento")
        if nasc:
            iso = _dia(nasc)
            if not iso or iso > date.today().isoformat() or iso < "1900-01-01":
                aviso(aba_cli, n, f"Data de nascimento \"{nasc}\" não entendida: ficou em branco.")
            else:
                dados["data_nascimento"] = iso
        uf = (l.get("estado") or "").strip().upper()
        if uf:
            if len(uf) == 2 and uf.isalpha():
                dados["estado"] = uf
            else:
                aviso(aba_cli, n, f"UF \"{uf}\" deve ter 2 letras: ficou em branco.")
        papeis = PAPEL_IMOB.get(_chave(l.get("papel"))) if segmento == "imobiliaria" and l.get("papel") else None
        tel = _digitos(dados.get("telefone"))
        if l.get("id"):
            atual = p_id.get(l["id"])
            if not atual:
                erro(aba_cli, n, "ID SAX não encontrado nesta empresa. Apague o ID para cadastrar como novo.")
                continue
            completo = True
        else:
            atual = achar_pessoa(doc, tel, email, nome)
            completo = False
        if doc and p_doc.get(doc) and atual is not p_doc.get(doc):
            aviso(aba_cli, n, "Este CPF/CNPJ já é de outro cadastro: ficou em branco.")
            dados.pop("cpf_cnpj", None)
            doc = ""
        if atual:
            mudar = {k: v for k, v in dados.items() if v and (completo or not atual.get(k)) and atual.get(k) != v}
            if papeis:
                mudar["papeis"] = sorted(set(atual.get("papeis") or []) | set(papeis))
            elif papel_cli not in (atual.get("papeis") or []) and "proprietario" not in (atual.get("papeis") or []):
                mudar["papeis"] = sorted(set(atual.get("papeis") or []) | {papel_cli})
            if "telefone" in mudar:
                mudar["telefone_digitos"] = tel or None
            if "cpf_cnpj" in mudar:
                mudar["documento_digitos"] = doc or None
            if mudar:
                cont["clientes_atualizados"] += 1
                if not simular and atual["id"] not in novos_ids:
                    atualizar.append(banco.pessoas.update_one({"id": atual["id"]}, {"$set": mudar}))
                atual.update(mudar)
            continue
        try:
            nova_pessoa({**{k: v for k, v in dados.items() if v is not None}, "papeis": papeis or [papel_cli]})
        except Exception as e:  # validação do modelo (tamanho de campo etc.)
            erro(aba_cli, n, f"Dados inválidos: {str(e).splitlines()[-1][:160]}")
            continue
        cont["clientes_novos"] += 1

    def temp(aba: str, n: int, valor: str, origem: str) -> str | None:
        if valor:
            t = normalizar_temp(valor)
            if t:
                return t
            aviso(aba, n, f"Temperatura \"{valor}\" não é Quente, Morno ou Frio: usei a regra da origem.")
        return pela_origem(origem, regras)

    def item_id(aba: str, n: int, codigo: str) -> str | None:
        if not codigo or not item:
            return None
        i = cod_item.get(_chave(codigo))
        if not i:
            aviso(aba, n, f"Código \"{codigo}\" não encontrado no estoque: ficou sem vínculo.")
        return i

    # ---- leads (caixa de entrada)
    novas_entradas: list[dict] = []
    if "leads" in tabelas:
        from models.crm import Entrada

        entradas = await banco.entradas.find({}, {"id": 1, "telefone": 1, "status": 1, "nome": 1, "email": 1, "origem": 1,
                                                  "temperatura": 1, "corretor_id": 1, "valor_estimado": 1}).to_list(None)
        e_id = {e["id"]: e for e in entradas}
        e_tel = {_digitos(e.get("telefone"))[-11:]: e for e in entradas if len(_digitos(e.get("telefone"))) >= 10 and e.get("status") != "convertido"}
        for n, l in tabelas["leads"]:
            nome = l.get("nome", "")
            if not nome:
                erro("LEADS", n, "Falta o nome.")
                continue
            tel = _digitos(l.get("telefone"))
            origem = (l.get("origem") or "Planilha")[:80]
            situacao = SITUACAO_LEAD.get(_chave(l.get("situacao")), None) if l.get("situacao") else None
            if l.get("situacao") and not situacao:
                aviso("LEADS", n, f"Situação \"{l['situacao']}\" não entendida: ficou Novo.")
            dados = {"nome": nome[:200], "telefone": l.get("telefone") or None, "email": (l.get("email") or "").lower() or None,
                     "origem": origem, "temperatura": temp("LEADS", n, l.get("temperatura", ""), origem),
                     "mensagem": (l.get("mensagem") or "")[:2000] or None, "valor_estimado": _valor(l.get("valor")),
                     "corretor_id": responsavel("LEADS", n, l.get("responsavel", ""))}
            if l.get("valor") and dados["valor_estimado"] is None:
                aviso("LEADS", n, f"Valor \"{l['valor']}\" não entendido: ficou em branco.")
            if segmento == "imobiliaria" and l.get("interesse"):
                dados["interesse"] = INTERESSE.get(_chave(l["interesse"]).split(" ")[0], "outro")
            it = item_id("LEADS", n, l.get("item", ""))
            if it:
                dados["imovel_id" if item == "imovel" else "veiculo_id"] = it
            if l.get("id"):
                atual = e_id.get(l["id"])
                if not atual:
                    erro("LEADS", n, "ID SAX não encontrado nesta empresa. Apague o ID para cadastrar como novo.")
                    continue
            else:
                atual = e_tel.get(tel[-11:]) if len(tel) >= 10 else None
            if atual:
                mudar = {k: v for k, v in dados.items() if v is not None and (l.get("id") or not atual.get(k)) and atual.get(k) != v}
                if situacao and situacao != "convertido" and atual.get("status") != "convertido":
                    mudar["status"] = situacao
                if mudar:
                    mudar["updated_at"] = agora
                    cont["leads_atualizados"] += 1
                    if not simular:
                        atualizar.append(banco.entradas.update_one({"id": atual["id"]}, {"$set": mudar}))
                continue
            criado = _dia(l.get("criado"))
            try:
                ent = Entrada(**{k: v for k, v in dados.items() if v is not None},
                              status=situacao if situacao in ("novo", "em_contato", "descartado") else "novo").model_dump()
            except Exception as e:
                erro("LEADS", n, f"Dados inválidos: {str(e).splitlines()[-1][:160]}")
                continue
            if ent.get("corretor_id"):
                ent["atribuido_em"] = agora
            if criado:
                ent["created_at"] = datetime.fromisoformat(criado).replace(hour=12, tzinfo=timezone.utc)
            ent["importado"] = True
            novas_entradas.append(ent)
            if len(tel) >= 10:
                e_tel[tel[-11:]] = ent
            cont["leads_novos"] += 1

    # ---- negócios
    novos_neg: list[dict] = []
    hist: list[dict] = []
    if "negocios" in tabelas:
        if not funil_padrao:
            raise ErroPlanilha("Esta empresa não tem funil. Crie um funil no CRM antes de importar negócios.")
        negs = {d["id"]: d for d in await banco.leads.find({}, {"id": 1, "funil_id": 1, "etapa_id": 1, "status": 1}).to_list(None)}
        por_nome_funil = {_chave(f["nome"]): f for f in funis}
        for n, l in tabelas["negocios"]:
            titulo = l.get("nome", "")
            if not titulo:
                erro("NEGOCIOS", n, "Falta o título.")
                continue
            f = funil_padrao
            if l.get("funil"):
                f = por_nome_funil.get(_chave(l["funil"]))
                if not f:
                    erro("NEGOCIOS", n, f"Funil \"{l['funil']}\" não existe. Use: {', '.join(x['nome'] for x in funis)}.")
                    continue
            etapas = f.get("etapas") or []
            etapa = etapas[0] if etapas else None
            if l.get("etapa"):
                etapa = next((e for e in etapas if _chave(e["nome"]) == _chave(l["etapa"])), None)
                if not etapa:
                    erro("NEGOCIOS", n, f"Etapa \"{l['etapa']}\" não existe no funil {f['nome']}.")
                    continue
            if not etapa:
                erro("NEGOCIOS", n, f"O funil {f['nome']} não tem etapas.")
                continue
            status = "aberto"
            if l.get("status"):
                status = STATUS_NEG.get(_chave(l["status"]))
                if not status:
                    erro("NEGOCIOS", n, f"Situação \"{l['status']}\" não entendida. Use Em andamento, Ganho ou Perdido.")
                    continue
            origem = (l.get("origem") or "Planilha")[:80]
            valor = _valor(l.get("valor"))
            if l.get("valor") and valor is None:
                aviso("NEGOCIOS", n, f"Valor \"{l['valor']}\" não entendido: ficou em branco.")
            previsao = _dia(l.get("previsao")) if l.get("previsao") else None
            if l.get("previsao") and not previsao:
                aviso("NEGOCIOS", n, f"Previsão \"{l['previsao']}\" não entendida: ficou em branco.")
            # cliente: acha pelo telefone, e-mail ou nome; senão cadastra
            cliente_id = None
            if l.get("cliente") or l.get("cliente_telefone") or l.get("cliente_email"):
                tel = _digitos(l.get("cliente_telefone"))
                mail = (l.get("cliente_email") or "").lower()
                pessoa = achar_pessoa("", tel, mail, l.get("cliente") or "")
                if not pessoa:
                    try:
                        pessoa = nova_pessoa({"nome": (l.get("cliente") or titulo)[:1000], "telefone": l.get("cliente_telefone") or None,
                                              "email": mail if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", mail or "") else None,
                                              "papeis": [papel_cli]})
                        cont["clientes_novos"] += 1
                    except Exception:
                        pessoa = None
                        aviso("NEGOCIOS", n, "Não deu para cadastrar o cliente desta linha: o negócio ficou sem cliente.")
                cliente_id = (pessoa or {}).get("id")
            dados = {"nome": titulo[:1000], "valor_estimado": valor, "origem": origem, "previsao_fechamento": previsao,
                     "temperatura": temp("NEGOCIOS", n, l.get("temperatura", ""), origem),
                     "corretor_id": responsavel("NEGOCIOS", n, l.get("responsavel", "")), "cliente_id": cliente_id,
                     "observacoes": (l.get("observacoes") or "")[:4000] or None,
                     "motivo_perda": (l.get("motivo_perda") or "")[:200] or None if status == "perdido" else None,
                     "etiquetas": [t.strip()[:40] for t in (l.get("etiquetas") or "").split(",") if t.strip()][:10]}
            it = item_id("NEGOCIOS", n, l.get("item", ""))
            if it:
                dados["imovel_id" if item == "imovel" else "veiculo_id"] = it
            if l.get("id"):
                atual = negs.get(l["id"])
                if not atual:
                    erro("NEGOCIOS", n, "ID SAX não encontrado nesta empresa. Apague o ID para cadastrar como novo.")
                    continue
                mudar = {k: v for k, v in dados.items() if v not in (None, [])}
                mudar.update(funil_id=f["id"], etapa_id=etapa["id"], status=status, estagio=estagio_legado(f, etapa["id"], status),
                             updated_at=agora)
                if status != atual.get("status"):
                    mudar["closed_at"] = agora if status != "aberto" else None
                if etapa["id"] != atual.get("etapa_id"):
                    mudar["etapa_desde"] = agora
                cont["negocios_atualizados"] += 1
                if not simular:
                    atualizar.append(banco.leads.update_one({"id": atual["id"]}, {"$set": mudar, "$addToSet": {"etapas_alcancadas": etapa["id"]}}))
                continue
            pos = next(i for i, e in enumerate(etapas) if e["id"] == etapa["id"])
            criado = _dia(l.get("criado"))
            criado_em = datetime.fromisoformat(criado).replace(hour=12, tzinfo=timezone.utc) if criado else agora
            doc = {"id": new_id(), **{k: v for k, v in dados.items() if v is not None}, "funil_id": f["id"], "etapa_id": etapa["id"],
                   "status": status, "estagio": estagio_legado(f, etapa["id"], status), "etapa_desde": criado_em,
                   "etapas_alcancadas": [e["id"] for e in etapas[:pos + 1]], "entrada_id": None,
                   "created_at": criado_em, "updated_at": agora, "closed_at": agora if status != "aberto" else None, "importado": True}
            doc.setdefault("etiquetas", [])
            novos_neg.append(doc)
            hist.append({"id": new_id(), "negocio_id": doc["id"], "tipo": "criado", "texto": "Negócio importado da planilha de dados",
                         "autor": por, "autor_id": por_usuario_id, "fixado": False, "em": agora})
            cont["negocios_novos"] += 1

    if not simular:
        from lib.importar_revendas import _em_paralelo

        if novas_pessoas:
            await banco.pessoas.insert_many(novas_pessoas)
        if novas_entradas:
            await banco.entradas.insert_many(novas_entradas)
        if novos_neg:
            await banco.leads.insert_many(novos_neg)
        if hist:
            await banco.historico.insert_many(hist)
        if atualizar:
            await _em_paralelo(atualizar)
    return {"simulado": simular, "contagens": cont, "erros": erros, "avisos": avisos,
            "abas": sorted(tabelas), "mais_mensagens": cont["linhas_com_erro"] > len(erros) or len(avisos) >= MAX_MENSAGENS}
